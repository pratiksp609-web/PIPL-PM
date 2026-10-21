"""Excel and PDF export helpers."""
from __future__ import annotations

import datetime as dt
import io

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


def _safe(v):
    if isinstance(v, str) and v[:1] in ("=", "+", "-", "@"):
        return "'" + v
    return v


def xlsx_bytes(sheets: list[tuple[str, list[str], list[list]]]) -> bytes:
    wb = Workbook()
    wb.remove(wb.active)
    for title, header, rows in sheets:
        ws = wb.create_sheet(title[:31])
        ws.append(header)
        for c in ws[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F3A5F")
            c.alignment = Alignment(vertical="center")
        for row in rows:
            ws.append([_safe(v) for v in row])
        for i, h in enumerate(header, 1):
            width = max([len(str(h))] + [len(str(r[i - 1])) for r in rows[:200] if i - 1 < len(r) and r[i - 1] is not None])
            ws.column_dimensions[get_column_letter(i)].width = min(max(width + 2, 10), 60)
        ws.freeze_panes = "A2"
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def report_pdf(report: dict) -> bytes:
    """Render a generated status report as a PDF using reportlab."""
    from reportlab.lib import colors
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
    from reportlab.lib.units import mm
    from reportlab.platypus import PageBreak, Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

    styles = getSampleStyleSheet()
    h1, h2, h3, body = styles["Title"], styles["Heading2"], styles["Heading3"], styles["BodyText"]
    small = ParagraphStyle("small", parent=body, fontSize=8, textColor=colors.HexColor("#555555"))
    buf = io.BytesIO()
    doc = SimpleDocTemplate(buf, pagesize=A4, leftMargin=15 * mm, rightMargin=15 * mm, topMargin=15 * mm, bottomMargin=15 * mm,
                            title=report["title"])
    story = []

    def p(text, style=body):
        story.append(Paragraph(_esc(text), style))

    def table(rows, widths=None):
        if not rows:
            return
        t = Table([[Paragraph(_esc(c), small) for c in r] for r in rows], colWidths=widths, repeatRows=1)
        t.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#1F3A5F")),
            ("TEXTCOLOR", (0, 0), (-1, 0), colors.white),
            ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#BBBBBB")),
            ("VALIGN", (0, 0), (-1, -1), "TOP"),
        ]))
        story.append(t)
        story.append(Spacer(1, 6))

    data = report["data"]
    p(report["title"], h1)
    p(f"Period: {data['period']['start']} to {data['period']['end']} · Generated {report['created_at']}", small)
    story.append(Spacer(1, 8))

    ai = report.get("ai_analysis")
    p("Analysis", h2)
    if report.get("ai_status") == "success" and ai:
        p("AI-assisted analysis based only on the records below. Facts, estimates and recommendations are labelled.", small)
        p(f"Overall health: {ai.get('overall_health', '-')}")
        p(ai.get("summary", ""))
        for label, key in (("Confirmed facts", "confirmed_facts"), ("Estimates & assumptions", "estimates"),
                           ("Risks", "risks"), ("Data issues", "data_issues"), ("Recommendations", "recommendations")):
            items = ai.get(key) or []
            if items:
                p(label, h3)
                for it in items:
                    p(f"• {it if isinstance(it, str) else it.get('text', '')}")
    else:
        p(f"AI analysis unavailable ({report.get('ai_status')}). Template-based summary:", small)
    for line in data.get("template_summary", []):
        p(f"• {line}")

    s = data["summary"]
    p("Company summary" if data["scope"] == "company" else "Project summary", h2)
    table([["Metric", "Value"]] + [[k.replace("_", " ").capitalize(), str(v)] for k, v in s.items() if not isinstance(v, (dict, list))])
    if data.get("financial"):
        p("Financial status (client receivables: X − Y = Z)", h2)
        rows = [["Currency", "Contract value (X)", "Verified received (Y)", "Outstanding (Z)", "Credit", "Overdue pending"]]
        for cur, v in data["financial"]["totals_by_currency"].items():
            rows.append([cur, v["contract_value"], v["verified_received"], v["outstanding"], v["credit"], v["overdue_pending"]])
        table(rows)

    for proj in data["projects"]:
        story.append(PageBreak() if len(data["projects"]) > 1 else Spacer(1, 6))
        p(f"{proj['code']} · {proj['name']}", h2)
        p(f"Client: {proj['client']} · Manager: {proj.get('manager') or '-'} · Status: {proj['status']} · Health: {proj['health']}", small)
        planned = proj.get("planned_progress")
        p(f"Progress: {proj['progress']}% actual" + (f" vs {planned}% planned (linear estimate)" if planned is not None else ""))
        if proj.get("receivable"):
            r = proj["receivable"]
            p(f"X {r['contract_value']} − Y {r['verified_received']} = Z {r['outstanding']} {r['currency']}"
              + (f" (credit {r['credit']})" if r["credit"] != "0.00" else ""))
        for title, key, cols in (
            ("Completed tasks in period", "completed_tasks", ("title", "completed_at")),
            ("Overdue tasks", "overdue_tasks", ("title", "due_date", "assignee")),
            ("Upcoming deadlines", "upcoming", ("title", "date", "kind")),
            ("Updates in period", "updates", ("update_at", "author", "progress", "work_completed")),
            ("Blockers & risks", "blockers", ("date", "text")),
        ):
            rows = proj.get(key) or []
            if rows:
                p(title, h3)
                table([[c.replace("_", " ").capitalize() for c in cols]] + [[str(r.get(c, "") or "")[:300] for c in cols] for r in rows[:25]])
        proc = proj.get("procurement")
        if proc:
            p("Procurement", h3)
            p(f"Estimated {proc['estimated_cost']} · Committed {proc['committed_cost']} · Supplier paid {proc['supplier_paid']} · "
              f"Pending items {proc['pending_procurement_count']} · Late deliveries {len(proc['late_deliveries'])}")
        if proj.get("next_actions"):
            p("Recommended next actions (rule-based)", h3)
            for a in proj["next_actions"]:
                p(f"• {a}")
    doc.build(story)
    return buf.getvalue()


def _esc(text) -> str:
    return (str(text) if text is not None else "").replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;").replace("\n", "<br/>")


def stamp() -> str:
    return dt.datetime.now().strftime("%Y%m%d-%H%M")
