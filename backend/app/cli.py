"""Management commands.

  python -m app.cli create-admin --email you@company.com --username admin --name "Your Name"
  python -m app.cli seed-sample        # development only: adds clearly marked [SAMPLE] records
  python -m app.cli reset-sample       # removes every sample record (and sample files)
  python -m app.cli run-checks         # run reminder / notification checks once
"""
from __future__ import annotations

import argparse
import datetime as dt
import getpass
import io
import os
import sys
from decimal import Decimal

from sqlalchemy import delete, or_, select, update

from . import models as m
from .config import get_settings
from .db import SessionLocal
from .finance import recalc_supplier_po
from .security import hash_password, validate_password_strength
from .storage import extract_commercial_fields, extract_text, save_bytes, storage_root

settings = get_settings()
SAMPLE_DOMAIN = "sample.pipl.local"
SAMPLE_PASSWORD = "Sample@1234"


def create_admin(email: str, username: str, name: str, password: str | None) -> None:
    password = password or os.environ.get("ADMIN_PASSWORD") or getpass.getpass("Admin password: ")
    validate_password_strength(password)
    db = SessionLocal()
    try:
        if db.scalar(select(m.User).where(or_(m.User.email == email.lower(), m.User.username == username))):
            print("A user with that email or username already exists.")
            sys.exit(1)
        db.add(m.User(email=email.lower(), username=username, full_name=name, role="admin", password_hash=hash_password(password)))
        db.commit()
        print(f"Admin user {email} created.")
    finally:
        db.close()


# ---------------------------------------------------------------- sample documents
def _quotation_docx(number: str, client: str, items: list[tuple[str, int, str, int]], date: dt.date, validity_days: int) -> bytes:
    import docx

    d = docx.Document()
    d.add_heading("PIPL Engineering — Commercial Quotation [SAMPLE]", 1)
    d.add_paragraph(f"Quotation No: {number}")
    d.add_paragraph(f"Date: {date:%d/%m/%Y}")
    d.add_paragraph(f"Client: {client}")
    d.add_paragraph(f"Validity: {validity_days} days from quotation date")
    t = d.add_table(rows=1, cols=5)
    for c, h in zip(t.rows[0].cells, ["Description", "Qty", "Unit", "Rate", "Amount"]):
        c.text = h
    subtotal = 0
    for desc, qty, unit, rate in items:
        row = t.add_row().cells
        row[0].text, row[1].text, row[2].text, row[3].text, row[4].text = desc, str(qty), unit, f"{rate:,}", f"{qty * rate:,}"
        subtotal += qty * rate
    gst = round(subtotal * 0.18)
    d.add_paragraph(f"Sub Total: Rs. {subtotal:,}.00")
    d.add_paragraph(f"GST @ 18%: Rs. {gst:,}.00")
    d.add_paragraph(f"Grand Total: Rs. {subtotal + gst:,}.00")
    d.add_paragraph("Payment Terms: 30% advance, 60% against delivery, 10% after commissioning")
    d.add_paragraph("Delivery Terms: Ex-works Pune, 10 weeks from PO")
    buf = io.BytesIO()
    d.save(buf)
    return buf.getvalue()


def _po_pdf(number: str, client: str, date: dt.date, subtotal: int, tax: int, terms: str) -> bytes:
    from reportlab.lib.pagesizes import A4
    from reportlab.pdfgen import canvas

    buf = io.BytesIO()
    c = canvas.Canvas(buf, pagesize=A4)
    y = 800
    for line in [f"{client} — PURCHASE ORDER [SAMPLE]", f"PO No: {number}", f"PO Date: {date:%d/%m/%Y}",
                 "Supplier: PIPL Engineering Pvt Ltd", f"Sub Total: Rs. {subtotal:,}.00", f"IGST 18%: Rs. {tax:,}.00",
                 f"Total Amount: Rs. {subtotal + tax:,}.00", f"Payment Terms: {terms}", "Delivery Terms: Ex-works Pune, 12 weeks from PO"]:
        c.drawString(60, y, line)
        y -= 24
    c.save()
    return buf.getvalue()


def _doc(db, user, project, category, filename, data, **links) -> m.Document:
    ext = filename.rsplit(".", 1)[-1]
    stored, digest = save_bytes(ext, data)
    text, _ = extract_text(ext, data)
    from .storage import MIME_BY_EXT

    d = m.Document(title=filename.rsplit(".", 1)[0], category=category, original_filename=filename, stored_filename=stored,
                   mime_type=MIME_BY_EXT[ext], extension=ext, size_bytes=len(data), sha256=digest, extracted_text=text,
                   project_id=project.id, client_id=project.client_id, uploaded_by_id=user.id, description="[SAMPLE] generated demo file", **links)
    db.add(d)
    db.flush()
    return d


def seed_sample() -> None:
    if settings.is_production or not settings.allow_sample_data:
        print("Refusing to load sample data in production.")
        sys.exit(1)
    db = SessionLocal()
    try:
        if db.scalar(select(m.Project).where(m.Project.is_sample.is_(True)).limit(1)):
            print("Sample data already loaded. Run reset-sample first to reload it.")
            return
        today = dt.date.today()
        now = dt.datetime.now(dt.timezone.utc)

        def user(username, name, role):
            u = db.scalar(select(m.User).where(m.User.username == username))
            if not u:
                u = m.User(email=f"{username}@{SAMPLE_DOMAIN}", username=username, full_name=f"{name} [SAMPLE]", role=role,
                           password_hash=hash_password(SAMPLE_PASSWORD))
                db.add(u)
                db.flush()
            return u

        admin = user("sample.admin", "Asha Admin", "admin")
        pm1 = user("sample.pm", "Rahul Mehta", "project_manager")
        pm2 = user("sample.pm2", "Neha Kulkarni", "project_manager")
        tm1 = user("sample.engineer", "Vikram Patil", "team_member")
        tm2 = user("sample.designer", "Priya Shah", "team_member")

        clients = []
        for name, company, contact, gst in [
            ("Tata Process Plants", "[SAMPLE] Tata Process Plants Ltd", "S. Iyer", "27AABCT1234F1Z5"),
            ("Kirloskar Foundry", "[SAMPLE] Kirloskar Foundry Works", "M. Deshpande", "27AACCK5678G1Z2"),
            ("Gujarat Agro Exports", "[SAMPLE] Gujarat Agro Exports Pvt Ltd", "R. Patel", "24AADCG9012H1Z8"),
        ]:
            c = m.Client(name=f"[SAMPLE] {name}", company_name=company, contact_person=contact, email=f"{contact.split()[-1].lower()}@{SAMPLE_DOMAIN}",
                         phone="+91 20 5555 0101", billing_address="Sample Industrial Estate, Pune, Maharashtra", gstin=gst,
                         is_sample=True, created_by_id=admin.id)
            db.add(c)
            clients.append(c)
        db.flush()

        specs = [
            ("Effluent Treatment Plant – 200 KLD", clients[0], pm1, today - dt.timedelta(days=120), today + dt.timedelta(days=60), "active", 55, "high", [tm1, tm2]),
            ("Sand Reclamation System Retrofit", clients[1], pm2, today - dt.timedelta(days=150), today - dt.timedelta(days=10), "active", 80, "critical", [tm1]),
            ("Cold Storage Turnkey Facility", clients[2], pm1, today + dt.timedelta(days=15), today + dt.timedelta(days=210), "planning", 0, "medium", [tm2]),
            ("Conveyor Line Automation", clients[0], pm2, today - dt.timedelta(days=300), today - dt.timedelta(days=40), "completed", 100, "medium", [tm1]),
        ]
        projects = []
        for i, (name, client, mgr, start, end, status, prog, prio, members) in enumerate(specs, 1):
            p = m.Project(code=f"SAMPLE-{today.year}-{i:03d}", name=f"[SAMPLE] {name}", client_id=client.id, manager_id=mgr.id,
                          description=f"Sample project for local development: {name}.", start_date=start, expected_completion_date=end,
                          actual_completion_date=end if status == "completed" else None, budget=Decimal("4000000") * i,
                          currency="INR", status=status, progress=prog, priority=prio, is_sample=True, created_by_id=admin.id)
            db.add(p)
            db.flush()
            for mem in members:
                db.add(m.ProjectMember(project_id=p.id, user_id=mem.id))
            projects.append(p)
        etp, sand, cold, conveyor = projects

        # quotations: ETP has R0 superseded and R1 approved; sand approved; cold submitted (expiring soon)
        def quotation(p, number, revisions):
            q = m.Quotation(project_id=p.id, client_id=p.client_id, quotation_number=number, title=f"Offer for {p.name}", created_by_id=pm1.id)
            db.add(q)
            db.flush()
            versions = []
            for rev, (items, status, qdate, vdays) in enumerate(revisions):
                data = _quotation_docx(f"{number} Rev {rev}", p.client.company_name, items, qdate, vdays)
                text, _ = extract_text("docx", data)
                extraction = extract_commercial_fields(text, "quotation", "docx", data)
                doc = _doc(db, pm1, p, "quotation", f"{number}-R{rev}.docx", data)
                subtotal = sum(q_ * r for _, q_, _, r in items)
                tax = round(subtotal * 0.18)
                v = m.QuotationVersion(
                    quotation_id=q.id, revision_number=rev, quotation_date=qdate, validity_date=qdate + dt.timedelta(days=vdays),
                    subtotal=subtotal, tax_amount=tax, total_amount=subtotal + tax, currency="INR",
                    payment_terms="30% advance, 60% against delivery, 10% after commissioning", delivery_terms="Ex-works Pune, 10 weeks from PO",
                    delivery_date=qdate + dt.timedelta(days=70), scope=f"Design, supply, installation and commissioning of {p.name}",
                    line_items=[{"description": d_, "quantity": str(q_), "unit": u_, "unit_price": str(r), "amount": str(q_ * r)} for d_, q_, u_, r in items],
                    status=status, document_id=doc.id, extraction=extraction, extraction_verified=True,
                    approved_by_id=pm1.id if status == "approved" else None, approved_at=now if status == "approved" else None,
                    created_by_id=pm1.id)
                db.add(v)
                db.flush()
                doc.quotation_version_id = v.id
                versions.append(v)
            return q, versions

        etp_items = [("Equalisation tank with mixers", 1, "set", 850000), ("MBBR bioreactor package", 1, "set", 1450000),
                     ("Clarifier and sludge handling", 1, "set", 620000), ("Electrical & instrumentation", 1, "lot", 380000)]
        etp_items_r1 = etp_items[:3] + [("Electrical & instrumentation", 1, "lot", 420000)]
        quotation(etp, "PIPL/Q/ETP/001", [(etp_items, "superseded", today - dt.timedelta(days=140), 30),
                                          (etp_items_r1, "approved", today - dt.timedelta(days=128), 45)])
        sand_items = [("Thermal reclamation unit", 1, "set", 2200000), ("Pneumatic conveying system", 1, "lot", 760000)]
        quotation(sand, "PIPL/Q/SRS/014", [(sand_items, "approved", today - dt.timedelta(days=160), 30)])
        cold_items = [("Insulated PUF panels", 1800, "sqm", 2100), ("Refrigeration plant 2 x 60 TR", 1, "set", 4200000)]
        quotation(cold, "PIPL/Q/CSF/022", [(cold_items, "submitted", today - dt.timedelta(days=25), 30)])
        conv_items = [("Belt conveyor 40 m", 2, "nos", 640000), ("PLC automation panel", 1, "set", 520000)]
        quotation(conveyor, "PIPL/Q/CLA/007", [(conv_items, "approved", today - dt.timedelta(days=310), 30)])

        def po(p, number, subtotal, tax, status, issue, terms, items, delivery_days):
            data = _po_pdf(number, p.client.company_name, issue, subtotal, tax, terms)
            doc = _doc(db, pm1, p, "client_po", f"{number}.pdf", data)
            text, _ = extract_text("pdf", data)
            x = m.ClientPurchaseOrder(project_id=p.id, client_id=p.client_id, po_number=number, issue_date=issue, po_value=subtotal,
                                      tax_amount=tax, total_value=subtotal + tax, currency="INR", payment_terms=terms,
                                      delivery_terms="Ex-works Pune, 12 weeks from PO", delivery_date=issue + dt.timedelta(days=delivery_days),
                                      scope=f"Design, supply, installation and commissioning of {p.name}", line_items=items, status=status,
                                      document_id=doc.id, extraction=extract_commercial_fields(text, "po", "pdf", data), extraction_verified=True,
                                      verified_by_id=pm1.id if status == "verified" else None, verified_at=now if status == "verified" else None,
                                      review_notes="[SAMPLE] verified against approved quotation" if status == "verified" else None,
                                      created_by_id=pm1.id)
            db.add(x)
            db.flush()
            doc.client_po_id = x.id
            return x

        def li(items):
            return [{"description": d_, "quantity": str(q_), "unit": u_, "unit_price": str(r), "amount": str(q_ * r)} for d_, q_, u_, r in items]

        etp_sub = sum(q_ * r for _, q_, _, r in etp_items_r1)
        etp_po = po(etp, "TPP/PO/2026/4471", etp_sub, round(etp_sub * 0.18), "verified", today - dt.timedelta(days=118),
                    "30% advance, 60% against delivery, 10% after commissioning", li(etp_items_r1), 70)
        sand_sub = sum(q_ * r for _, q_, _, r in sand_items)
        po(sand, "KFW/PO/889", sand_sub - 60000, round((sand_sub - 60000) * 0.18), "under_review", today - dt.timedelta(days=150),
           "20% advance, 70% against delivery, 10% after commissioning",
           li([("Thermal reclamation unit", 1, "set", 2200000), ("Pneumatic conveying system", 1, "lot", 700000)]), 84)
        conv_sub = sum(q_ * r for _, q_, _, r in conv_items)
        conv_po = po(conveyor, "TPP/PO/2025/3120", conv_sub, round(conv_sub * 0.18), "verified", today - dt.timedelta(days=300),
                     "50% advance, 50% on completion", li(conv_items), 200)

        etp_total = etp_po.total_value
        pays = [
            (etp, etp_po, etp_total * Decimal("0.30"), today - dt.timedelta(days=110), None, "advance", "verified", "SAMPLE-UTR-1001"),
            (etp, etp_po, Decimal("1200000"), today - dt.timedelta(days=20), None, "milestone", "verified", "SAMPLE-UTR-1002"),
            (etp, etp_po, Decimal("250000"), today - dt.timedelta(days=3), None, "partial", "recorded", "SAMPLE-UTR-1003"),
            (etp, etp_po, Decimal("900000"), None, today - dt.timedelta(days=5), "milestone", "pending", None),
            (etp, etp_po, Decimal("500000"), None, today + dt.timedelta(days=25), "milestone", "pending", None),
            (conveyor, conv_po, conv_po.total_value * Decimal("0.5"), today - dt.timedelta(days=290), None, "advance", "verified", "SAMPLE-UTR-0901"),
            (conveyor, conv_po, conv_po.total_value * Decimal("0.5") + 15000, today - dt.timedelta(days=35), None, "final", "verified", "SAMPLE-UTR-0902"),
        ]
        for p, cpo, amount, paid_on, due, ptype, status, ref in pays:
            db.add(m.PaymentReceived(project_id=p.id, client_id=p.client_id, client_po_id=cpo.id, amount=amount.quantize(Decimal("0.01")),
                                     payment_date=paid_on, due_date=due, payment_type=ptype, milestone_label=f"{ptype.capitalize()} [SAMPLE]",
                                     method="bank_transfer" if paid_on else None, transaction_reference=ref, invoice_reference=f"SAMPLE-INV-{ref or 'TBD'}",
                                     status=status, verified_by_id=pm1.id if status == "verified" else None, created_by_id=pm1.id,
                                     notes="[SAMPLE] demo payment" + (" (includes 15,000 overpayment)" if ptype == "final" else "")))

        suppliers = []
        for name in ("[SAMPLE] Shree Steel Fabricators", "[SAMPLE] Aqua Pumps & Motors", "[SAMPLE] Delta Controls"):
            s = m.Supplier(name=name, contact_person="Sales Desk", email=f"sales@{SAMPLE_DOMAIN}", phone="+91 22 5555 0199", is_sample=True)
            db.add(s)
            suppliers.append(s)
        db.flush()

        bom_specs = [
            (etp, "MS equalisation tank 50 m³", "IS 2062 E250, 6 mm shell", 1, "nos", 420000, suppliers[0], -30),
            (etp, "Submersible mixers 3 kW", "SS316 impeller", 2, "nos", 145000, suppliers[1], -10),
            (etp, "MBBR media", "PP, 500 m²/m³", 18, "m³", 22000, None, 15),
            (etp, "PLC & SCADA panel", "Siemens S7-1200", 1, "set", 360000, suppliers[2], 25),
            (sand, "Rotary drum lining", "Refractory castable", 1, "lot", 380000, suppliers[0], -45),
            (sand, "Blower 15 kW", "Twin lobe", 2, "nos", 190000, suppliers[1], -20),
        ]
        boms = []
        for p, name, spec, qty, unit, rate, sup, days in bom_specs:
            b = m.BOMItem(project_id=p.id, item_name=name, specification=spec, quantity=qty, unit=unit, estimated_rate=rate,
                          supplier_id=sup.id if sup else None, required_date=today + dt.timedelta(days=days), procurement_status="pending",
                          remarks="[SAMPLE]", created_by_id=pm1.id)
            db.add(b)
            boms.append(b)
        db.flush()

        def spo(p, sup, number, po_date, delivery, status, lines):
            x = m.SupplierPurchaseOrder(project_id=p.id, supplier_id=sup.id, po_number=number, po_date=po_date, delivery_date=delivery,
                                        status=status, payment_terms="30 days credit", notes="[SAMPLE]", created_by_id=pm1.id)
            x.items = [m.SupplierPOItem(bom_item_id=b.id if b else None, description=d_, quantity=q_, unit=u_, unit_price=pr, tax_rate=18,
                                        received_quantity=rq, amount=0) for b, d_, q_, u_, pr, rq in lines]
            recalc_supplier_po(x)
            db.add(x)
            db.flush()
            return x

        s1 = spo(etp, suppliers[0], "PIPL/SPO/1001", today - dt.timedelta(days=80), today - dt.timedelta(days=35), "received",
                 [(boms[0], boms[0].item_name, 1, "nos", 435000, 1)])
        s2 = spo(etp, suppliers[1], "PIPL/SPO/1002", today - dt.timedelta(days=40), today - dt.timedelta(days=6), "ordered",
                 [(boms[1], boms[1].item_name, 2, "nos", 139000, 0)])
        s3 = spo(sand, suppliers[1], "PIPL/SPO/1003", today - dt.timedelta(days=70), today - dt.timedelta(days=25), "partially_received",
                 [(boms[5], boms[5].item_name, 2, "nos", 198000, 1)])
        boms[0].procurement_status, boms[0].actual_rate, boms[0].received_quantity = "received", 435000, 1
        boms[1].procurement_status, boms[1].actual_rate = "delayed", 139000
        boms[5].procurement_status, boms[5].actual_rate, boms[5].received_quantity = "partially_received", 198000, 1
        d = _po_pdf("PIPL/SPO/1001", "PIPL Engineering", s1.po_date, 435000, 78300, "30 days credit")
        s1.document_id = _doc(db, pm1, etp, "supplier_po", "PIPL-SPO-1001.pdf", d, supplier_po_id=s1.id, supplier_id=suppliers[0].id).id
        db.add(m.SupplierPayment(supplier_po_id=s1.id, supplier_id=suppliers[0].id, project_id=etp.id, amount=Decimal("300000"),
                                 payment_date=today - dt.timedelta(days=30), method="bank_transfer", reference="SAMPLE-SUP-PAY-01",
                                 status="verified", notes="[SAMPLE] supplier payment — not part of client receivables"))
        _ = (s2, s3)

        task_specs = [
            (etp, "Finalize MBBR media supplier", tm1, "high", "in_progress", 40, -3),
            (etp, "Civil foundation inspection", tm2, "medium", "todo", 0, 4),
            (etp, "Submit P&ID revision C for approval", tm2, "high", "under_review", 90, 1),
            (etp, "Mixer delivery follow-up", tm1, "critical", "blocked", 20, -1),
            (sand, "Commissioning plan", tm1, "high", "in_progress", 60, -12),
            (sand, "Operator training material", tm1, "low", "todo", 0, 9),
            (cold, "Site survey and soil report", tm2, "medium", "todo", 0, 20),
            (conveyor, "Handover documentation", tm1, "medium", "completed", 100, -45),
        ]
        for p, title, who, prio, status, prog, due in task_specs:
            db.add(m.Task(project_id=p.id, title=f"{title} [SAMPLE]", assignee_id=who.id, priority=prio, status=status, progress=prog,
                          due_date=today + dt.timedelta(days=due), start_date=today - dt.timedelta(days=20),
                          completed_at=now - dt.timedelta(days=46) if status == "completed" else None, created_by_id=p.manager_id))

        for p, author, days_ago, prog, health, done, blockers in [
            (etp, tm1, 9, 48, "on_track", "Equalisation tank erected and hydro-tested.", None),
            (etp, tm2, 2, 55, "at_risk", "MBBR reactor shell fabrication 70% complete; cable trays installed.", "Mixer delivery delayed by supplier."),
            (sand, tm1, 6, 80, "delayed", "Reclamation unit installed; trial run pending blower.", "Second blower not yet received."),
        ]:
            db.add(m.ProjectUpdate(project_id=p.id, author_id=author.id, update_at=now - dt.timedelta(days=days_ago), health=health,
                                   work_completed=f"{done} [SAMPLE]", progress=prog, work_in_progress="Electrical works",
                                   blockers=blockers, next_steps="Continue as per schedule", remarks="[SAMPLE] demo update"))

        db.add(m.CalendarEvent(title="[SAMPLE] Client review meeting – ETP", event_type="meeting", project_id=etp.id, client_id=etp.client_id,
                               start_at=dt.datetime.combine(today + dt.timedelta(days=2), dt.time(5, 30), dt.timezone.utc),
                               end_at=dt.datetime.combine(today + dt.timedelta(days=2), dt.time(6, 30), dt.timezone.utc),
                               location="Client site, Pune", remind_minutes_before=60, created_by_id=pm1.id))
        db.add(m.CalendarEvent(title="[SAMPLE] Factory acceptance test", event_type="milestone", project_id=sand.id,
                               start_at=dt.datetime.combine(today + dt.timedelta(days=6), dt.time(4, 0), dt.timezone.utc), all_day=True,
                               created_by_id=pm2.id))
        db.commit()
        print("Sample data loaded. All records are prefixed [SAMPLE].")
        print(f"Sample logins (password {SAMPLE_PASSWORD}): sample.admin, sample.pm, sample.pm2, sample.engineer, sample.designer")
    finally:
        db.close()


def reset_sample() -> None:
    if settings.is_production:
        print("Refusing to modify data in production.")
        sys.exit(1)
    db = SessionLocal()
    try:
        pids = list(db.scalars(select(m.Project.id).where(m.Project.is_sample.is_(True))))
        uids = list(db.scalars(select(m.User.id).where(m.User.email.like(f"%@{SAMPLE_DOMAIN}"))))
        files = list(db.scalars(select(m.Document.stored_filename).where(m.Document.project_id.in_(pids))))
        if pids:
            db.execute(update(m.Document).where(m.Document.project_id.in_(pids)).values(
                quotation_version_id=None, client_po_id=None, supplier_po_id=None, payment_id=None, supplier_payment_id=None,
                task_id=None, update_id=None, report_id=None, superseded_by_id=None))
            for model in (m.QuotationVersion, m.ClientPurchaseOrder, m.SupplierPurchaseOrder):
                sub = select(m.Quotation.id).where(m.Quotation.project_id.in_(pids))
                cond = model.quotation_id.in_(sub) if model is m.QuotationVersion else model.project_id.in_(pids)
                db.execute(update(model).where(cond).values(document_id=None))
            db.execute(delete(m.SupplierPayment).where(m.SupplierPayment.project_id.in_(pids)))
            db.execute(delete(m.PaymentReceived).where(m.PaymentReceived.project_id.in_(pids)))
            db.execute(delete(m.SupplierPOItem).where(m.SupplierPOItem.supplier_po_id.in_(
                select(m.SupplierPurchaseOrder.id).where(m.SupplierPurchaseOrder.project_id.in_(pids)))))
            db.execute(delete(m.SupplierPurchaseOrder).where(m.SupplierPurchaseOrder.project_id.in_(pids)))
            db.execute(delete(m.BOMItem).where(m.BOMItem.project_id.in_(pids)))
            db.execute(delete(m.ClientPurchaseOrder).where(m.ClientPurchaseOrder.project_id.in_(pids)))
            db.execute(delete(m.QuotationVersion).where(m.QuotationVersion.quotation_id.in_(select(m.Quotation.id).where(m.Quotation.project_id.in_(pids)))))
            db.execute(delete(m.Quotation).where(m.Quotation.project_id.in_(pids)))
            db.execute(delete(m.Document).where(m.Document.project_id.in_(pids)))
            db.execute(delete(m.ProjectUpdateRevision).where(m.ProjectUpdateRevision.update_id.in_(select(m.ProjectUpdate.id).where(m.ProjectUpdate.project_id.in_(pids)))))
            db.execute(delete(m.ProjectUpdate).where(m.ProjectUpdate.project_id.in_(pids)))
            db.execute(delete(m.CalendarEvent).where(m.CalendarEvent.project_id.in_(pids)))
            db.execute(delete(m.Task).where(m.Task.project_id.in_(pids)))
            db.execute(update(m.GeneratedReport).where(m.GeneratedReport.project_id.in_(pids)).values(project_id=None))
            db.execute(delete(m.Project).where(m.Project.id.in_(pids)))
        db.execute(delete(m.Client).where(m.Client.is_sample.is_(True)))
        db.execute(delete(m.Supplier).where(m.Supplier.is_sample.is_(True)))
        if uids:
            db.execute(delete(m.Task).where(or_(m.Task.created_by_id.in_(uids), m.Task.assignee_id.in_(uids)), m.Task.project_id.is_(None)))
            db.execute(delete(m.GeneratedReport).where(m.GeneratedReport.generated_by_id.in_(uids)))
            db.execute(delete(m.User).where(m.User.id.in_(uids)))
        db.commit()
        root = storage_root()
        for f in files:
            (root / f).unlink(missing_ok=True)
        print(f"Removed sample data: {len(pids)} projects, {len(uids)} users, {len(files)} files.")
    finally:
        db.close()


def main(argv: list[str] | None = None) -> None:
    parser = argparse.ArgumentParser(prog="python -m app.cli")
    sub = parser.add_subparsers(dest="cmd", required=True)
    a = sub.add_parser("create-admin")
    a.add_argument("--email", required=True)
    a.add_argument("--username", required=True)
    a.add_argument("--name", required=True)
    a.add_argument("--password", help="omit to be prompted (or set ADMIN_PASSWORD)")
    sub.add_parser("seed-sample")
    sub.add_parser("reset-sample")
    sub.add_parser("run-checks")
    args = parser.parse_args(argv)
    if args.cmd == "create-admin":
        create_admin(args.email, args.username, args.name, args.password)
    elif args.cmd == "seed-sample":
        seed_sample()
    elif args.cmd == "reset-sample":
        reset_sample()
    elif args.cmd == "run-checks":
        from .scheduler import run_all_checks

        print(run_all_checks())


if __name__ == "__main__":
    main()
