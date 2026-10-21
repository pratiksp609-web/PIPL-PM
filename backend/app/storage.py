"""Secure file storage, content validation, text extraction, field extraction and previews.

Files are stored outside any web root under random names and are only served through
authenticated endpoints with Content-Disposition and nosniff headers. Nothing uploaded is executed.
"""
from __future__ import annotations

import datetime as dt
import hashlib
import io
import logging
import re
import secrets
import zipfile
from decimal import Decimal, InvalidOperation
from pathlib import Path

from .config import get_settings
from .core import APIError

log = logging.getLogger("pipl.storage")
settings = get_settings()

MIME_BY_EXT = {
    "pdf": "application/pdf",
    "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    "doc": "application/msword",
    "xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    "xls": "application/vnd.ms-excel",
    "csv": "text/csv",
    "txt": "text/plain",
    "png": "image/png",
    "jpg": "image/jpeg",
    "jpeg": "image/jpeg",
}
CATEGORIES = ("quotation", "client_po", "supplier_po", "invoice", "payment_proof", "report", "drawing", "other")
OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"


def storage_root() -> Path:
    root = Path(settings.upload_dir)
    root.mkdir(parents=True, exist_ok=True)
    return root


def _sniff_ok(ext: str, data: bytes) -> bool:
    head = data[:8]
    if ext == "pdf":
        return data[:5] == b"%PDF-"
    if ext in ("docx", "xlsx"):
        if head[:4] != b"PK\x03\x04":
            return False
        try:
            names = zipfile.ZipFile(io.BytesIO(data)).namelist()
        except zipfile.BadZipFile:
            return False
        return ("word/document.xml" in names) if ext == "docx" else ("xl/workbook.xml" in names)
    if ext in ("doc", "xls"):
        return head == OLE_MAGIC
    if ext == "png":
        return head == b"\x89PNG\r\n\x1a\n"
    if ext in ("jpg", "jpeg"):
        return head[:3] == b"\xff\xd8\xff"
    if ext in ("txt", "csv"):
        if b"\x00" in data[:4096]:
            return False
        try:
            data[:65536].decode("utf-8")
        except UnicodeDecodeError:
            try:
                data[:65536].decode("latin-1")
            except UnicodeDecodeError:
                return False
        return True
    return False


def validate_upload(filename: str, data: bytes) -> tuple[str, str]:
    """Return (extension, mime). Raises APIError on any problem."""
    if not filename or "." not in filename:
        raise APIError(422, "File must have an extension", "invalid_file")
    ext = filename.rsplit(".", 1)[-1].lower()
    if ext not in MIME_BY_EXT:
        raise APIError(422, f"File type .{ext} is not allowed. Allowed: {', '.join(sorted(MIME_BY_EXT))}", "invalid_file_type")
    if not data:
        raise APIError(422, "The uploaded file is empty", "empty_file")
    if len(data) > settings.max_upload_mb * 1024 * 1024:
        raise APIError(413, f"File exceeds the {settings.max_upload_mb} MB limit", "file_too_large")
    if not _sniff_ok(ext, data):
        raise APIError(422, f"File content does not match the .{ext} format", "content_mismatch")
    return ext, MIME_BY_EXT[ext]


def save_bytes(ext: str, data: bytes) -> tuple[str, str]:
    stored = f"{secrets.token_hex(16)}.{ext}"
    path = storage_root() / stored
    with open(path, "xb") as fh:
        fh.write(data)
    return stored, hashlib.sha256(data).hexdigest()


def file_path(stored_filename: str) -> Path:
    if not re.fullmatch(r"[0-9a-f]{32}\.[a-z]{2,5}", stored_filename):
        raise APIError(400, "Invalid file reference", "invalid_file")
    path = storage_root() / stored_filename
    if not path.exists():
        raise APIError(404, "The stored file is missing from disk", "file_missing")
    return path


def safe_download_name(name: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9._ -]", "_", name).strip() or "download"
    return cleaned[:150]


# ---------------------------------------------------------------- text extraction
def extract_text(ext: str, data: bytes) -> tuple[str, str]:
    """Return (text, method). method describes how text was obtained or why it was not."""
    try:
        if ext == "pdf":
            from pypdf import PdfReader

            reader = PdfReader(io.BytesIO(data))
            text = "\n".join((page.extract_text() or "") for page in reader.pages[:50])
            if text.strip():
                return text, "pdf_text_layer"
            return "", "no_text_layer_ocr_unavailable" if not _ocr_available() else _ocr_pdf_unsupported()
        if ext == "docx":
            import docx

            d = docx.Document(io.BytesIO(data))
            parts = [p.text for p in d.paragraphs]
            for table in d.tables:
                for row in table.rows:
                    parts.append(" | ".join(c.text.strip() for c in row.cells))
            return "\n".join(parts), "docx"
        if ext == "xlsx":
            from openpyxl import load_workbook

            wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
            parts = []
            for ws in wb.worksheets[:5]:
                for row in ws.iter_rows(max_row=500, values_only=True):
                    vals = [str(v) for v in row if v is not None]
                    if vals:
                        parts.append(" | ".join(vals))
            return "\n".join(parts), "xlsx"
        if ext in ("txt", "csv"):
            try:
                return data.decode("utf-8"), "text"
            except UnicodeDecodeError:
                return data.decode("latin-1"), "text"
        if ext in ("png", "jpg", "jpeg"):
            if _ocr_available():
                import pytesseract
                from PIL import Image

                return pytesseract.image_to_string(Image.open(io.BytesIO(data))), "ocr"
            return "", "image_ocr_unavailable"
    except Exception as exc:  # parsing libraries raise many exception types on malformed files
        log.warning("text extraction failed", extra={"ext": ext, "error": str(exc)})
        return "", "extraction_failed"
    return "", "unsupported_format"


def _ocr_available() -> bool:
    try:
        import shutil

        import pytesseract  # noqa: F401

        return shutil.which("tesseract") is not None
    except ImportError:
        return False


def _ocr_pdf_unsupported() -> str:
    return "no_text_layer_scanned_pdf"


# ---------------------------------------------------------------- structured field extraction
_AMOUNT = r"(?:rs\.?|inr|₹|usd|\$|eur|€)?\s*([0-9]{1,3}(?:,[0-9]{2,3})+(?:\.[0-9]{1,2})?|[0-9]+(?:\.[0-9]{1,2})?)"
_DATE_PATTERNS = [
    (r"\b(\d{4})-(\d{1,2})-(\d{1,2})\b", "ymd"),
    (r"\b(\d{1,2})[/.\-](\d{1,2})[/.\-](\d{4})\b", "dmy"),
    (r"\b(\d{1,2})(?:st|nd|rd|th)?\s+([A-Za-z]{3,9})[,]?\s+(\d{4})\b", "d_mon_y"),
    (r"\b([A-Za-z]{3,9})\s+(\d{1,2})(?:st|nd|rd|th)?,?\s+(\d{4})\b", "mon_d_y"),
]
_MONTHS = {m: i for i, m in enumerate(["jan", "feb", "mar", "apr", "may", "jun", "jul", "aug", "sep", "oct", "nov", "dec"], 1)}


def _parse_date(s: str) -> dt.date | None:
    for pat, kind in _DATE_PATTERNS:
        mt = re.search(pat, s)
        if not mt:
            continue
        try:
            a, b, c = mt.groups()
            if kind == "ymd":
                return dt.date(int(a), int(b), int(c))
            if kind == "dmy":  # Indian convention: day first
                return dt.date(int(c), int(b), int(a))
            if kind == "d_mon_y":
                return dt.date(int(c), _MONTHS[b[:3].lower()], int(a))
            if kind == "mon_d_y":
                return dt.date(int(c), _MONTHS[a[:3].lower()], int(b))
        except (ValueError, KeyError):
            continue
    return None


def _money(s: str) -> Decimal | None:
    try:
        return Decimal(s.replace(",", ""))
    except (InvalidOperation, AttributeError):
        return None


def _field(value, confidence: str, source: str | None) -> dict:
    v = value.isoformat() if isinstance(value, dt.date) else (str(value) if isinstance(value, Decimal) else value)
    return {"value": v, "confidence": confidence, "source": (source or "")[:160]}


def _labelled(text: str, labels: str, value_pat: str):
    # Label and value must be on the same line ([ \t] instead of \s), so a heading such as "QUOTATION" cannot
    # swallow the next line's label as its value.
    pat = re.compile(rf"(?:{labels})[ \t]*(?:no\.?|number|#|amount|value)?[ \t]*[:\-–]?[ \t]*{value_pat}", re.IGNORECASE)
    mt = pat.search(text)
    return (mt.group(1), mt.group(0)) if mt else (None, None)


def _labelled_amount(text: str, labels: str, exclude: str | None = None) -> tuple[Decimal | None, str | None]:
    """Find an amount on the first line matching a label. `labels` is tried alternative by alternative, in priority
    order, so "grand total" wins over a plain "total"; lines matching `exclude` (e.g. sub-total) are skipped."""
    lines = text.splitlines()
    for label in labels.split("|"):
        for line in lines:
            if exclude and re.search(exclude, line, re.IGNORECASE):
                continue
            if re.search(rf"\b(?:{label})\b", line, re.IGNORECASE):
                amounts = re.findall(_AMOUNT, line, re.IGNORECASE)
                vals = [_money(a) for a in amounts if _money(a) is not None]
                vals = [v for v in vals if v >= 1]
                if vals:
                    return vals[-1], line.strip()
    return None, None


def _docx_line_items(data: bytes) -> list[dict]:
    import docx

    items = []
    d = docx.Document(io.BytesIO(data))
    for table in d.tables:
        if len(table.rows) < 2:
            continue
        headers = [c.text.strip().lower() for c in table.rows[0].cells]

        def col(*names):
            for i, h in enumerate(headers):
                if any(n in h for n in names):
                    return i
            return None

        di, qi, ri, ai, ui = col("description", "item", "particular"), col("qty", "quantity"), col("rate", "unit price", "price"), col("amount", "total", "value"), col("unit", "uom")
        if di is None or (qi is None and ai is None):
            continue
        for row in table.rows[1:]:
            cells = [c.text.strip() for c in row.cells]
            desc = cells[di] if di < len(cells) else ""
            if not desc or re.search(r"\b(sub\s*total|total|gst|tax)\b", desc, re.IGNORECASE):
                continue
            item = {"description": desc[:500]}
            for key, idx in (("quantity", qi), ("unit_price", ri), ("amount", ai)):
                if idx is not None and idx < len(cells):
                    v = _money(re.sub(r"[^0-9.,]", "", cells[idx]))
                    if v is not None:
                        item[key] = str(v)
            if ui is not None and ui < len(cells) and ui not in (qi, ri, ai, di):
                item["unit"] = cells[ui][:20]
            items.append(item)
    return items[:200]


def extract_commercial_fields(text: str, kind: str, ext: str | None = None, data: bytes | None = None) -> dict:
    """Heuristically extract quotation / PO fields. Every field carries a confidence level and must be verified."""
    fields: dict[str, dict] = {}
    if not text.strip():
        return {"fields": fields, "line_items": [], "warnings": ["No readable text found; enter details manually."]}

    if kind == "quotation":
        num, src = _labelled(text, r"quotation|quote|offer|proposal|ref(?:erence)?", r"((?=[A-Z0-9/\-_.]*\d)[A-Z0-9][A-Z0-9/\-_.]{2,40})")
    else:
        num, src = _labelled(text, r"p\.?\s?o\.?|purchase\s+order|order", r"((?=[A-Z0-9/\-_.]*\d)[A-Z0-9][A-Z0-9/\-_.]{2,40})")
    if num:
        fields["number"] = _field(num, "medium", src)

    rev, src = _labelled(text, r"rev(?:ision)?\.?", r"([0-9]{1,3})\b")
    if rev and kind == "quotation":
        fields["revision_number"] = _field(int(rev), "medium", src)

    for line in text.splitlines():
        low = line.lower()
        if "date" in low and "valid" not in low and "delivery" not in low and "date" not in fields:
            d = _parse_date(line)
            if d:
                fields["date"] = _field(d, "medium", line)
        if re.search(r"valid(?:ity)?\s*(?:till|until|upto|up to|date)", low):
            d = _parse_date(line)
            if d:
                fields["validity_date"] = _field(d, "medium", line)
        if "delivery" in low and "delivery_date" not in fields:
            d = _parse_date(line)
            if d:
                fields["delivery_date"] = _field(d, "low", line)
    if "date" not in fields:
        d = _parse_date(text)
        if d:
            fields["date"] = _field(d, "low", "first date found in document")
    if "validity_date" not in fields and "date" in fields:
        mt = re.search(r"valid(?:ity)?[^\n]{0,30}?(\d{1,3})\s*days", text, re.IGNORECASE)
        if mt:
            base = dt.date.fromisoformat(fields["date"]["value"])
            fields["validity_date"] = _field(base + dt.timedelta(days=int(mt.group(1))), "low", mt.group(0))

    total, tsrc = _labelled_amount(text, r"grand\s*total|total\s*amount|net\s*amount|total\s*value|po\s*value|total",
                                   exclude=r"sub\s*-?\s*total")
    subtotal, ssrc = _labelled_amount(text, r"sub\s*-?\s*total|basic\s*(?:value|amount)|taxable\s*value")
    tax_total = Decimal(0)
    tax_src = []
    for line in text.splitlines():
        if re.search(r"\b(igst|cgst|sgst|gst|vat|tax)\b", line, re.IGNORECASE) and not re.search(r"gstin|total", line, re.IGNORECASE):
            amounts = [_money(a) for a in re.findall(_AMOUNT, line, re.IGNORECASE)]
            amounts = [a for a in amounts if a is not None and a >= 1]
            # ignore pure percentages such as "18%"
            if amounts and not re.search(rf"{re.escape(str(amounts[-1]))}\s*%", line):
                tax_total += amounts[-1]
                tax_src.append(line.strip())
    consistent = subtotal is not None and total is not None and tax_total and abs(subtotal + tax_total - total) <= Decimal("1.00")
    if total is not None:
        fields["total_amount"] = _field(total, "high" if consistent else "medium", tsrc)
    if subtotal is not None:
        fields["subtotal"] = _field(subtotal, "high" if consistent else "medium", ssrc)
    if tax_total:
        fields["tax_amount"] = _field(tax_total, "high" if consistent else "low", "; ".join(tax_src))
    elif subtotal is not None and total is not None and total >= subtotal:
        fields["tax_amount"] = _field(total - subtotal, "low", "derived as total - subtotal")

    if re.search(r"₹|\binr\b|\brs\.?\b", text, re.IGNORECASE):
        fields["currency"] = _field("INR", "medium", "currency symbol")
    elif re.search(r"\busd\b|\$", text, re.IGNORECASE):
        fields["currency"] = _field("USD", "medium", "currency symbol")
    elif re.search(r"\beur\b|€", text, re.IGNORECASE):
        fields["currency"] = _field("EUR", "medium", "currency symbol")

    for key, label in (("payment_terms", r"payment\s*terms?"), ("delivery_terms", r"delivery\s*terms?|delivery\s*period|delivery")):
        mt = re.search(rf"(?:{label})\s*[:\-–]\s*(.+)", text, re.IGNORECASE)
        if mt:
            fields[key] = _field(mt.group(1).strip()[:500], "medium", mt.group(0))

    line_items = []
    if ext == "docx" and data:
        try:
            line_items = _docx_line_items(data)
        except Exception:  # malformed tables should not block upload
            line_items = []

    warnings = ["Automatically extracted values may be wrong. Review every field before saving."]
    if total is not None and subtotal is not None and not consistent:
        warnings.append("Subtotal + taxes does not match the total; check the amounts.")
    return {"fields": fields, "line_items": line_items, "warnings": warnings}


# ---------------------------------------------------------------- previews
def docx_preview_html(data: bytes) -> str:
    import mammoth
    import nh3

    result = mammoth.convert_to_html(io.BytesIO(data))
    return nh3.clean(
        result.value,
        tags={"p", "h1", "h2", "h3", "h4", "h5", "h6", "strong", "em", "u", "ul", "ol", "li", "table", "thead", "tbody",
              "tr", "td", "th", "br", "span", "sup", "sub", "blockquote"},
        attributes={"td": {"colspan", "rowspan"}, "th": {"colspan", "rowspan"}},
    )


def xlsx_preview_rows(data: bytes, max_rows: int = 200) -> list[list[str]]:
    from openpyxl import load_workbook

    wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    ws = wb.worksheets[0]
    rows = []
    for row in ws.iter_rows(max_row=max_rows, values_only=True):
        rows.append(["" if v is None else str(v) for v in row][:30])
    return rows
